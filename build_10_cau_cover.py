# -*- coding: utf-8 -*-
"""
CHƯƠNG TRÌNH XÂY DỰNG TÀI LIỆU WORD: BÁO CÁO TRẢ LỜI CHI TIẾT 10 CÂU HỎI
MÔN: XỬ LÝ DỮ LIỆU LỚN (CTKU122) - TRƯỜNG ĐH TÀI NGUYÊN VÀ MÔI TRƯỜNG HÀ NỘI
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from build_10_cau_helpers import (
    set_cell_background, set_cell_margins, add_header_paragraph,
    add_body_paragraph, add_bullet_point, add_callout_box,
    add_code_snippet, add_image_with_caption, format_custom_table
)

def build_cover_page(doc):
    # Header School
    p_sch = doc.add_paragraph()
    p_sch.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sch.paragraph_format.space_before = Pt(0)
    p_sch.paragraph_format.space_after = Pt(2)
    r1 = p_sch.add_run("BỘ TÀI NGUYÊN VÀ MÔI TRƯỜNG\n")
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(11)
    r1.font.bold = True
    
    r2 = p_sch.add_run("TRƯỜNG ĐẠI HỌC TÀI NGUYÊN VÀ MÔI TRƯỜNG HÀ NỘI\n")
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(13)
    r2.font.bold = True
    r2.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
    
    r3 = p_sch.add_run("KHOA CÔNG NGHỆ THÔNG TIN - BỘ MÔN KHOA HỌC MÁY TÍNH\n")
    r3.font.name = "Times New Roman"
    r3.font.size = Pt(11.5)
    r3.font.bold = True
    r3.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)
    
    r4 = p_sch.add_run("-----------------------***-----------------------")
    r4.font.name = "Times New Roman"
    r4.font.size = Pt(10)
    
    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(36)
    p_title.paragraph_format.space_after = Pt(12)
    
    rt1 = p_title.add_run("BÁO CÁO BÀI TẬP LỚN\n")
    rt1.font.name = "Times New Roman"
    rt1.font.size = Pt(18)
    rt1.font.bold = True
    rt1.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
    
    rt2 = p_title.add_run("HỌC PHẦN: XỬ LÝ DỮ LIỆU LỚN (CTKU122)\n\n")
    rt2.font.name = "Times New Roman"
    rt2.font.size = Pt(15)
    rt2.font.bold = True
    rt2.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)
    
    rt3 = p_title.add_run("BẢN TRẢ LỜI CHI TIẾT 10 CÂU HỎI THỰC HÀNH & ĐÁNH GIÁ ĐIỂM (10,0 ĐIỂM)\n")
    rt3.font.name = "Times New Roman"
    rt3.font.size = Pt(14)
    rt3.font.bold = True
    rt3.font.color.rgb = RGBColor(0xC5, 0x30, 0x30)
    
    rt4 = p_title.add_run("ĐỀ TÀI: XÂY DỰNG HỆ THỐNG PHÂN TÍCH VÀ GIÁM SÁT CHẤT LƯỢNG MẠNG VIỄN THÔNG (TELECOM QUALITY OF SERVICE ANALYTICS) TRÊN NỀN TẢNG HADOOP, SPARK VÀ R")
    rt4.font.name = "Times New Roman"
    rt4.font.size = Pt(13)
    rt4.font.bold = True
    rt4.font.italic = True
    rt4.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
    
    doc.add_paragraph().paragraph_format.space_before = Pt(20)
    
    # Table Info
    tbl = doc.add_table(rows=1, cols=3)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.columns[0].width = Inches(2.2)
    tbl.columns[1].width = Inches(3.0)
    tbl.columns[2].width = Inches(1.3)
    
    headers = ["Họ và tên sinh viên", "Mã sinh viên & Lớp", "Tỷ lệ đóng góp"]
    rows = [
        ["Hoàng Ngọc Thi (Nhóm trưởng)", "21111062000 - ĐH13C2", "25%"],
        ["Nguyễn Viết Hoàng Anh", "21111062001 - ĐH13C2", "25%"],
        ["Nguyễn Việt Hòa", "21111062002 - ĐH13C2", "25%"],
        ["Nguyễn Quốc Anh", "21111062003 - ĐH13C2", "25%"]
    ]
    format_custom_table(tbl, [Inches(2.5), Inches(2.5), Inches(1.5)], headers, rows)
    
    p_gv = doc.add_paragraph()
    p_gv.paragraph_format.space_before = Pt(24)
    p_gv.paragraph_format.space_after = Pt(4)
    p_gv.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    rg1 = p_gv.add_run("Giảng viên hướng dẫn: ThS. Đỗ Thị Hồng Nhung\n")
    rg1.font.name = "Times New Roman"
    rg1.font.size = Pt(12)
    rg1.font.bold = True
    
    rg2 = p_gv.add_run("Bộ môn: Khoa học máy tính - Khoa Công nghệ thông tin\n")
    rg2.font.name = "Times New Roman"
    rg2.font.size = Pt(11.5)
    rg2.font.italic = True
    
    p_date = doc.add_paragraph()
    p_date.paragraph_format.space_before = Pt(40)
    p_date.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rd = p_date.add_run("Hà Nội, Năm 2026")
    rd.font.name = "Times New Roman"
    rd.font.size = Pt(12)
    rd.font.bold = True
    
    doc.add_page_break()

print("Cover page builder ready.")
