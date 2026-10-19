# -*- coding: utf-8 -*-
"""
CHƯƠNG TRÌNH TẠO FILE WORD: TRẢ LỜI CHI TIẾT 10 CÂU HỎI BÀI TẬP LỚN
MÔN: XỬ LÝ DỮ LIỆU LỚN (CTKU122) - TRƯỜNG ĐH TN&MT HÀ NỘI (HUNRE)
ĐỀ TÀI: PHÂN TÍCH VÀ GIÁM SÁT CHẤT LƯỢNG MẠNG VIỄN THÔNG (TELECOM QOS ANALYTICS)
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_header_paragraph(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.3
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.bold = True
    if level == 1:
        run.font.size = Pt(15)
        run.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D) # Dark Navy
    elif level == 2:
        run.font.size = Pt(13.5)
        run.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0) # Blue
    else:
        run.font.size = Pt(12.5)
        run.font.italic = True
        run.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
    return p

def add_body_paragraph(doc, text, bold_prefix="", italic_prefix=""):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.35
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    if bold_prefix:
        r_b = p.add_run(bold_prefix)
        r_b.font.name = "Times New Roman"
        r_b.font.size = Pt(12)
        r_b.font.bold = True
        
    if italic_prefix:
        r_i = p.add_run(italic_prefix)
        r_i.font.name = "Times New Roman"
        r_i.font.size = Pt(12)
        r_i.font.italic = True
        
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    return p

def add_bullet_point(doc, bold_title, content):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    r_b = p.add_run(bold_title + ": ")
    r_b.font.name = "Times New Roman"
    r_b.font.size = Pt(12)
    r_b.font.bold = True
    r_b.font.color.rgb = RGBColor(0x1A, 0x20, 0x2C)
    
    r_c = p.add_run(content)
    r_c.font.name = "Times New Roman"
    r_c.font.size = Pt(12)
    return p

def add_callout_box(doc, title, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.5)
    
    cell = table.cell(0, 0)
    set_cell_background(cell, "F7FAFC")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Left border highlight
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="0" w:color="2B6CB0"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    
    r1 = p.add_run(f"ĐỀ BÀI YÊU CẦU - {title.upper()}\n")
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(11)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)
    
    r2 = p.add_run(text)
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(11.5)
    r2.font.italic = True
    r2.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)
    
    doc.add_paragraph().paragraph_format.space_before = Pt(4)

def add_code_snippet(doc, code_text, caption=""):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.5)
    
    cell = table.cell(0, 0)
    set_cell_background(cell, "2D3748") # Dark Slate
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    
    run = p.add_run(code_text.strip())
    run.font.name = "Consolas"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(0xE2, 0xE8, 0xF0)
    
    if caption:
        p_cap = doc.add_paragraph()
        p_cap.paragraph_format.space_before = Pt(3)
        p_cap.paragraph_format.space_after = Pt(6)
        p_cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rc = p_cap.add_run(f"Đoạn mã: {caption}")
        rc.font.name = "Times New Roman"
        rc.font.size = Pt(10.5)
        rc.font.italic = True
        rc.font.color.rgb = RGBColor(0x71, 0x80, 0x96)

def add_image_with_caption(doc, img_path, caption, width=Inches(6.0)):
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.paragraph_format.space_before = Pt(6)
        p_img.paragraph_format.space_after = Pt(3)
        p_img.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.add_run().add_picture(img_path, width=width)
        
        p_cap = doc.add_paragraph()
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(8)
        p_cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rc = p_cap.add_run(caption)
        rc.font.name = "Times New Roman"
        rc.font.size = Pt(10.5)
        rc.font.italic = True
        rc.font.bold = True
        rc.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)

def format_custom_table(table, col_widths, headers, data_rows):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    # Header Row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "1A365D")
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=120, right=120)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(11)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            
    # Data Rows
    for r_idx, row in enumerate(data_rows):
        row_cells = table.add_row().cells
        bg_color = "F7FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row):
            row_cells[c_idx].text = str(val)
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=80, bottom=80, left=100, right=100)
            p = row_cells[c_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx > 1 else WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(10.5)
                r.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
                
    # Apply widths
    for row in table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width

print("Builder helper functions initialized successfully.")
