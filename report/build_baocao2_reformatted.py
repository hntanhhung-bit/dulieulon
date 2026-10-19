# -*- coding: utf-8 -*-
"""
CHƯƠNG TRÌNH XÂY DỰNG BÁO CÁO BÀI TẬP LỚN XỬ LÝ DỮ LIỆU LỚN (CTKU122)
ĐỀ TÀI: PHÂN TÍCH CHẤT LƯỢNG MẠNG VIỄN THÔNG (TELECOM QUALITY OF SERVICE ANALYTICS)
TRƯỜNG ĐẠI HỌC TÀI NGUYÊN VÀ MÔI TRƯỜNG HÀ NỘI - KHOA CÔNG NGHỆ THÔNG TIN

Cấu trúc mục lục và đề mục chuẩn 100% theo mẫu thiết kế chi tiết:
- Chương 1: Tổng quan bài toán và công cụ xử lý dữ liệu lớn (1.1 -> 1.1.1..1.1.4, 1.2, 1.3 -> 1.3.1..1.3.11)
- Chương 2: Dữ liệu và quy trình tiền xử lý (2.1 -> 2.1.1..2.1.4, 2.2, 2.3)
- Chương 3: Xây dựng hệ thống xử lý dữ liệu lớn (3.1 -> 3.1.1..3.1.2, 3.2 -> 3.2.1..3.2.4, 3.3 -> 3.3.1..3.3.2, 3.4 -> 3.4.1..3.4.3, 3.5 -> 3.5.1..3.5.3, 3.6 -> 3.6.1..3.6.4)
- Chương 4: Kết quả và đánh giá (4.1, 4.2 -> 4.2.1..4.2.3, 4.3 -> 4.3.1..4.3.3, 4.4, 4.5)
- Chương 5: Kết luận và hướng phát triển (5.1, 5.2, 5.3)
- Kết luận, Tài liệu tham khảo, Phụ lục A, B, C
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

base_dir = r"c:\dulieulon"
img_dir = os.path.join(base_dir, "report")
template_path = os.path.join(base_dir, "Mau_Bao_cao_BTL_Xu_ly_du_lieu_lon_CTKU122 (1).docx")
output_path = os.path.join(base_dir, "BaoCao 2.docx")

# ==============================================================================
# HÀM BỔ TRỢ ĐỊNH DẠNG (STYLE & XML HELPERS)
# ==============================================================================
def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_p(doc, text="", bold_prefix=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=6, line_spacing=1.35, italic=False):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    if bold_prefix:
        r_b = p.add_run(bold_prefix)
        r_b.font.name = 'Times New Roman'
        r_b.font.size = Pt(13)
        r_b.font.bold = True
        r_b.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    if text:
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(13)
        r.font.italic = italic
        r.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
    return p

def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.30
    p.paragraph_format.left_indent = Inches(0.25)
    if bold_prefix:
        r_b = p.add_run(bold_prefix)
        r_b.font.name = 'Times New Roman'
        r_b.font.size = Pt(13)
        r_b.font.bold = True
        r_b.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
    return p

def add_h1(doc, text):
    p = doc.add_paragraph(style='Heading 1')
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.line_spacing = 1.2
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    return p

def add_h2(doc, text):
    p = doc.add_paragraph(style='Heading 2')
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.left_indent = Inches(0.0)
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    return p

def add_h3(doc, text):
    p = doc.add_paragraph(style='Heading 3')
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.left_indent = Inches(0.0)
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
    return p

def add_tbl(doc, headers, data, caption=None, col_widths=None):
    if caption:
        cp = doc.add_paragraph()
        cp.paragraph_format.space_before = Pt(12)
        cp.paragraph_format.space_after = Pt(4)
        cp.paragraph_format.keep_with_next = True
        c_run = cp.add_run(caption)
        c_run.font.name = 'Times New Roman'
        c_run.font.size = Pt(11.5)
        c_run.font.italic = True
        c_run.font.bold = True
        c_run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Tiêu đề bảng
    hdr_cells = table.rows[0].cells
    for idx, header_text in enumerate(headers):
        hdr_cells[idx].text = header_text
        set_cell_background(hdr_cells[idx], "1E3A8A")
        set_cell_margins(hdr_cells[idx], top=100, bottom=100, left=140, right=140)
        p = hdr_cells[idx].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        for r in p.runs:
            r.font.name = 'Times New Roman'
            r.font.size = Pt(10.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # Dữ liệu bảng
    for r_idx, row_data in enumerate(data):
        row_cells = table.rows[r_idx + 1].cells
        bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, cell_value in enumerate(row_data):
            row_cells[c_idx].text = str(cell_value)
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=80, bottom=80, left=120, right=120)
            p = row_cells[c_idx].paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            if c_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(10.5)

    if col_widths:
        for row in table.rows:
            for idx, w in enumerate(col_widths):
                row.cells[idx].width = Inches(w)

    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)
    return table

def add_fig(doc, img_name, caption, width_in=5.8):
    img_path = os.path.join(img_dir, img_name)
    if os.path.exists(img_path):
        actual_width = min(float(width_in), 5.85)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run()
        run.add_picture(img_path, width=Inches(actual_width))

        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.line_spacing = 1.15
        cp.paragraph_format.space_before = Pt(4)
        cp.paragraph_format.space_after = Pt(16)
        c_run = cp.add_run(caption)
        c_run.font.name = 'Times New Roman'
        c_run.font.size = Pt(11.5)
        c_run.font.italic = True
        c_run.font.bold = True
        c_run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

def add_code(doc, code_str, caption=None):
    if caption:
        cp = doc.add_paragraph()
        cp.paragraph_format.space_before = Pt(8)
        cp.paragraph_format.space_after = Pt(2)
        cp.paragraph_format.keep_with_next = True
        c_run = cp.add_run(caption)
        c_run.font.name = 'Times New Roman'
        c_run.font.size = Pt(11)
        c_run.font.italic = True
        c_run.font.bold = True

    tbl = doc.add_table(rows=1, cols=1)
    tbl.style = 'Table Grid'
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F1F5F9")
    set_cell_margins(cell, top=100, bottom=100, left=150, right=150)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(code_str)
    r.font.name = 'Consolas'
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)

def add_callout(doc, title, text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.style = 'Table Grid'
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "EFF6FF")
    set_cell_margins(cell, top=120, bottom=120, left=160, right=160)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.25
    r_t = p.add_run(f"📌 {title}: ")
    r_t.font.name = 'Times New Roman'
    r_t.font.size = Pt(11.5)
    r_t.font.bold = True
    r_t.font.color.rgb = RGBColor(0x1D, 0x4E, 0xD8)

    r_txt = p.add_run(text)
    r_txt.font.name = 'Times New Roman'
    r_txt.font.size = Pt(11)
    r_txt.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)

def add_toc_item(doc, title, page_str, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.25), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    
    if level == 1:
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Inches(0.0)
        r_t = p.add_run(title)
        r_t.font.name = 'Times New Roman'
        r_t.font.size = Pt(12)
        r_t.font.bold = True
        r_t.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        p.add_run("\t")
        r_p = p.add_run(str(page_str))
        r_p.font.name = 'Times New Roman'
        r_p.font.size = Pt(12)
        r_p.font.bold = True
        r_p.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    elif level == 2:
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.left_indent = Inches(0.20)
        r_t = p.add_run(title)
        r_t.font.name = 'Times New Roman'
        r_t.font.size = Pt(12)
        r_t.font.bold = True
        r_t.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
        p.add_run("\t")
        r_p = p.add_run(str(page_str))
        r_p.font.name = 'Times New Roman'
        r_p.font.size = Pt(12)
        r_p.font.bold = True
        r_p.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
    else:  # level 3
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.left_indent = Inches(0.45)
        r_t = p.add_run(title)
        r_t.font.name = 'Times New Roman'
        r_t.font.size = Pt(11.5)
        r_t.font.bold = False
        r_t.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
        p.add_run("\t")
        r_p = p.add_run(str(page_str))
        r_p.font.name = 'Times New Roman'
        r_p.font.size = Pt(11.5)
        r_p.font.bold = False
        r_p.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    return p

print("Loaded formatting helpers successfully.")

if __name__ == "__main__":
    if sys.stdout.encoding != 'utf-8':
        try:
            sys.stdout.reconfigure(encoding='utf-8')
        except Exception:
            pass

    current_dir = os.path.dirname(os.path.abspath(__file__))
    if current_dir not in sys.path:
        sys.path.insert(0, current_dir)

    print(f"Loading template: {template_path}")
    doc = docx.Document(template_path)

    # Import parts
    from report_sections_part1 import (
        build_frontmatter_and_toc,
        build_chapter_1,
        build_chapter_2
    )
    from report_sections_part2 import (
        build_chapter_3,
        build_chapter_4,
        build_chapter_5,
        build_conclusion_and_references
    )

    print("Executing: Frontmatter & Table of Contents...")
    build_frontmatter_and_toc(doc)

    print("Executing: Chapter 1...")
    build_chapter_1(doc)

    print("Executing: Chapter 2...")
    build_chapter_2(doc)

    print("Executing: Chapter 3...")
    build_chapter_3(doc)

    print("Executing: Chapter 4...")
    build_chapter_4(doc)

    print("Executing: Chapter 5...")
    build_chapter_5(doc)

    print("Executing: Conclusion & References & Appendices...")
    build_conclusion_and_references(doc)

    print(f"Saving final report to: {output_path}")
    try:
        doc.save(output_path)
        print(f"==> SUCCESS: Saved report to {output_path}")
    except PermissionError:
        alt_path = os.path.join(base_dir, "BaoCao 2_updated.docx")
        print(f"==> NOTICE: '{output_path}' is open in Word or locked. Saving to: {alt_path}")
        doc.save(alt_path)
        print(f"==> SUCCESS: Saved report to {alt_path}")
