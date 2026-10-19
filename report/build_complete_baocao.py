# ==============================================================================
# HỆ THỐNG XÂY DỰNG BÁO CÁO BÀI TẬP LỚN TOÀN DIỆN (CHUẨN HUNRE - CTKU122)
# ==============================================================================

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

template_path = r"c:\dulieulon\Mau_Bao_cao_BTL_Xu_ly_du_lieu_lon_CTKU122 (1).docx"
output_docx_path = r"c:\dulieulon\BaoCao.docx"
img_dir = r"c:\dulieulon\report"

doc = docx.Document(template_path)

# Cấu hình lề trang theo chuẩn: Trái 3cm, Phải 2cm, Trên 2cm, Dưới 2cm
for section in doc.sections:
    section.top_margin = Inches(2.0 / 2.54)
    section.bottom_margin = Inches(2.0 / 2.54)
    section.left_margin = Inches(3.0 / 2.54)
    section.right_margin = Inches(2.0 / 2.54)

# ------------------------------------------------------------------------------
# 1. CẬP NHẬT TRANG BÌA & BẢNG THÔNG TIN MẪU
# ------------------------------------------------------------------------------
# P6: ĐỀ TÀI SỐ
doc.paragraphs[6].text = "ĐỀ TÀI SỐ 01"
doc.paragraphs[6].runs[0].font.name = "Times New Roman"
doc.paragraphs[6].runs[0].font.size = Pt(14)
doc.paragraphs[6].runs[0].font.bold = True

# P7: TÊN ĐỀ TÀI
doc.paragraphs[7].text = "XÂY DỰNG HỆ THỐNG PHÂN TÍCH VÀ GIÁM SÁT CHẤT LƯỢNG MẠNG VIỄN THÔNG (TELECOM QUALITY OF SERVICE ANALYTICS) TRÊN NỀN TẢNG HADOOP, SPARK VÀ R"
doc.paragraphs[7].runs[0].font.name = "Times New Roman"
doc.paragraphs[7].runs[0].font.size = Pt(16)
doc.paragraphs[7].runs[0].font.bold = True
doc.paragraphs[7].alignment = WD_ALIGN_PARAGRAPH.CENTER

# P8: Lớp học phần
doc.paragraphs[8].text = "Lớp học phần: ĐH11.CNTT - Xử lý Dữ liệu lớn"
doc.paragraphs[8].runs[0].font.name = "Times New Roman"
doc.paragraphs[8].runs[0].font.size = Pt(13)

# P9: Giảng viên giảng dạy
doc.paragraphs[9].text = "Giảng viên giảng dạy: ThS. / TS. Giảng viên phụ trách học phần"
doc.paragraphs[9].runs[0].font.name = "Times New Roman"
doc.paragraphs[9].runs[0].font.size = Pt(13)

# Table 0: Cập nhật danh sách sinh viên thực hiện
sv_info = [
    ("1", "211101001", "Nguyễn Văn A (Nhóm trưởng)", "ĐH11.CNTT"),
    ("2", "211101002", "Trần Thị B", "ĐH11.CNTT"),
    ("3", "211101003", "Lê Hoàng C", "ĐH11.CNTT"),
]
for idx, (tt, msv, hoten, lop) in enumerate(sv_info):
    row_cells = doc.tables[0].rows[idx + 1].cells
    row_cells[0].text = tt
    row_cells[1].text = msv
    row_cells[2].text = hoten
    row_cells[3].text = lop
    for c in row_cells:
        for p in c.paragraphs:
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(11)

# Table 5: Phân công nhiệm vụ thành viên
nhiemvu_info = [
    ("1", "211101001", "Nguyễn Văn A", "Thiết kế kiến trúc hệ thống, cấu hình Flume Ingestion, lưu trữ HDFS và thuật toán MapReduce Java", "35%"),
    ("2", "211101002", "Trần Thị B", "Xây dựng pipeline Spark DataFrame, Spark SQL và kịch bản phân cụm dữ liệu với Apache Pig Latin", "35%"),
    ("3", "211101003", "Lê Hoàng C", "Thiết kế bảng HBase NoSQL, phân tích R & SQLite, mô hình Logistic Regression và soạn thảo báo cáo", "30%"),
]
for idx, (tt, msv, hoten, nv, tl) in enumerate(nhiemvu_info):
    row_cells = doc.tables[5].rows[idx + 1].cells
    row_cells[0].text = tt
    row_cells[1].text = msv
    row_cells[2].text = hoten
    row_cells[3].text = nv
    row_cells[4].text = tl
    for c in row_cells:
        for p in c.paragraphs:
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(10.5)

# Cắt bỏ các đoạn outline hướng dẫn sau Table 5 để điền nội dung toàn diện mới
body = doc._body._element
children = list(body)
cut_index = -1
tbl_count = 0
for i, el in enumerate(children):
    if el.tag.endswith('tbl'):
        tbl_count += 1
        if tbl_count == 6: # Sau Table 5
            cut_index = i
            break

for el in children[cut_index + 2:]:
    if el.tag.endswith('sectPr'):
        continue
    body.remove(el)

print("Front matter configured and placeholder outlines cleared.")
