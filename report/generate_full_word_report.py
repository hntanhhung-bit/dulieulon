# -*- coding: utf-8 -*-
"""
CHƯƠNG TRÌNH XÂY DỰNG TOÀN DIỆN BÁO CÁO BÀI TẬP LỚN XỬ LÝ DỮ LIỆU LỚN (CTKU122)
TRƯỜNG ĐẠI HỌC TÀI NGUYÊN VÀ MÔI TRƯỜNG HÀ NỘI - HUNRE
KHOA CÔNG NGHỆ THÔNG TIN - BỘ MÔN CÔNG NGHỆ THÔNG TIN

Đề tài: XÂY DỰNG HỆ THỐNG PHÂN TÍCH VÀ GIÁM SÁT CHẤT LƯỢNG MẠNG VIỄN THÔNG
(TELECOM QUALITY OF SERVICE ANALYTICS) TRÊN NỀN TẢNG HADOOP, SPARK VÀ R

Tuân thủ nghiêm ngặt các quy tắc định dạng:
• Khổ giấy A4; lề trái 3 cm, phải 2 cm, trên 2 cm, dưới 2 cm.
• Font Times New Roman; cỡ chữ nội dung 13; giãn dòng 1.35.
• Tên chương cỡ 16, in đậm; mục cấp 1 cỡ 14, in đậm; mục cấp 2 cỡ 13, in đậm nghiêng.
• Hình và bảng được đánh số theo chương, có tên và được nhận xét trong nội dung.
• Source code dài để ở phụ lục hoặc nộp kèm; nội dung chính chỉ trình bày đoạn mã cần thiết.
• Báo cáo khuyến nghị 25-40 trang, không tính phụ lục.
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Thêm thư mục report vào sys.path để import các module
script_dir = os.path.dirname(os.path.abspath(__file__))
if script_dir not in sys.path:
    sys.path.insert(0, script_dir)

from report_builders_p1 import build_toc_and_lists, build_introduction, build_chapter_1
from report_builders_p2 import build_chapter_2, build_chapter_3
from report_builders_p3 import build_chapter_4, build_chapter_5, build_references, build_appendices

template_path = r"c:\dulieulon\Mau_Bao_cao_BTL_Xu_ly_du_lieu_lon_CTKU122 (1).docx"
output_docx_path = r"c:\dulieulon\BaoCao.docx"

print(f"Loading template from: {template_path}")
doc = docx.Document(template_path)

# ==============================================================================
# 1. CẤU HÌNH LỀ TRANG THEO QUY CHUẨN A4: TRÁI 3CM, PHẢI 2CM, TRÊN 2CM, DƯỚI 2CM
# ==============================================================================
for section in doc.sections:
    section.top_margin = Inches(2.0 / 2.54)
    section.bottom_margin = Inches(2.0 / 2.54)
    section.left_margin = Inches(3.0 / 2.54)
    section.right_margin = Inches(2.0 / 2.54)

# ==============================================================================
# 2. CẬP NHẬT TRANG BÌA & BẢNG THÔNG TIN MẪU
# ==============================================================================
# Đề tài số
doc.paragraphs[6].text = "ĐỀ TÀI SỐ 01"
doc.paragraphs[6].runs[0].font.name = "Times New Roman"
doc.paragraphs[6].runs[0].font.size = Pt(14)
doc.paragraphs[6].runs[0].font.bold = True

# Tên đề tài
doc.paragraphs[7].text = "XÂY DỰNG HỆ THỐNG PHÂN TÍCH VÀ GIÁM SÁT CHẤT LƯỢNG MẠNG VIỄN THÔNG (TELECOM QUALITY OF SERVICE ANALYTICS) TRÊN NỀN TẢNG HADOOP, SPARK VÀ R"
doc.paragraphs[7].runs[0].font.name = "Times New Roman"
doc.paragraphs[7].runs[0].font.size = Pt(16)
doc.paragraphs[7].runs[0].font.bold = True
doc.paragraphs[7].alignment = WD_ALIGN_PARAGRAPH.CENTER

# Lớp học phần
doc.paragraphs[8].text = "Lớp học phần: ĐH13C2 - Xử lý Dữ liệu lớn"
doc.paragraphs[8].runs[0].font.name = "Times New Roman"
doc.paragraphs[8].runs[0].font.size = Pt(13)

# Giảng viên giảng dạy
doc.paragraphs[9].text = "Giảng viên giảng dạy: ThS. / TS. Giảng viên phụ trách học phần"
doc.paragraphs[9].runs[0].font.name = "Times New Roman"
doc.paragraphs[9].runs[0].font.size = Pt(13)

# Table 0: Cập nhật danh sách sinh viên thực hiện (Trang bìa)
sv_list = [
    ("1", "211101001", "Hoàng Ngọc Thi (Nhóm trưởng)", "ĐH13C2"),
    ("2", "211101002", "Nguyễn Viết Hoàng Anh", "ĐH13C2"),
    ("3", "211101003", "Nguyễn Việt Hòa", "ĐH13C2"),
    ("4", "211101004", "Nguyễn Quốc Anh", "ĐH13C2"),
]
if len(doc.tables[0].rows) > 5:
    doc.tables[0]._element.remove(doc.tables[0].rows[5]._tr)

for idx, (tt, msv, hoten, lop) in enumerate(sv_list):
    r_cells = doc.tables[0].rows[idx + 1].cells
    r_cells[0].text, r_cells[1].text, r_cells[2].text, r_cells[3].text = tt, msv, hoten, lop
    for c in r_cells:
        for p in c.paragraphs:
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(11)

# Table 3: Bảng tổng hợp điểm đánh giá học phần
score_sv_list = [
    ("1", "211101001", "Hoàng Ngọc Thi"),
    ("2", "211101002", "Nguyễn Viết Hoàng Anh"),
    ("3", "211101003", "Nguyễn Việt Hòa"),
    ("4", "211101004", "Nguyễn Quốc Anh"),
]
if len(doc.tables[3].rows) > 5:
    doc.tables[3]._element.remove(doc.tables[3].rows[5]._tr)

for idx, (tt, msv, hoten) in enumerate(score_sv_list):
    r_cells = doc.tables[3].rows[idx + 1].cells
    r_cells[0].text, r_cells[1].text, r_cells[2].text = tt, msv, hoten
    for c in r_cells:
        for p in c.paragraphs:
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(10.5)

# Table 5: Phân công nhiệm vụ thành viên - Chia đều 25% mỗi thành viên (Tổng 100%)
tasks_list = [
    ("1", "211101001", "Hoàng Ngọc Thi", "Quản trị dự án; thiết kế kiến trúc hệ thống Lambda Architecture; thiết lập hạ tầng HDFS và cấu hình Flume Ingestion", "25%"),
    ("2", "211101002", "Nguyễn Viết Hoàng Anh", "Thiết kế & lập trình thuật toán MapReduce Java (Custom Composite Key, Partitioner, Combiner); đóng gói JAR và benchmark YARN", "25%"),
    ("3", "211101003", "Nguyễn Việt Hòa", "Xây dựng pipeline Apache Spark DataFrame & Spark SQL (tối ưu cache/persist); phát triển kịch bản phân cụm dữ liệu với Apache Pig Latin", "25%"),
    ("4", "211101004", "Nguyễn Quốc Anh", "Thiết kế cơ sở dữ liệu NoSQL HBase độ trễ thấp; xây dựng mô hình Học máy Logistic Regression bằng R & SQLite; trực quan hóa và tổng hợp báo cáo", "25%"),
]
if len(doc.tables[5].rows) > 5:
    doc.tables[5]._element.remove(doc.tables[5].rows[5]._tr)

for idx, (tt, msv, hoten, nv, tl) in enumerate(tasks_list):
    r_cells = doc.tables[5].rows[idx + 1].cells
    r_cells[0].text, r_cells[1].text, r_cells[2].text, r_cells[3].text, r_cells[4].text = tt, msv, hoten, nv, tl
    for c in r_cells:
        for p in c.paragraphs:
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(10.5)

# Cắt bỏ các đoạn outline sau Table 5 để điền nội dung toàn diện mới
body = doc._body._element
children = list(body)
cut_index = -1
tbl_count = 0
for i, el in enumerate(children):
    if el.tag.endswith('tbl'):
        tbl_count += 1
        if tbl_count == 6:  # Sau Table 5
            cut_index = i
            break

for el in children[cut_index + 2:]:
    if el.tag.endswith('sectPr'):
        continue
    body.remove(el)

print("Front matter configured and placeholder outlines cleared.")

# ==============================================================================
# 3. XÂY DỰNG NỘI DUNG TỪNG PHẦN
# ==============================================================================
print("Building Table of Contents & Lists...")
build_toc_and_lists(doc)

print("Building Introduction...")
build_introduction(doc)

print("Building Chapter 1: Overview & Big Data Problem...")
build_chapter_1(doc)

print("Building Chapter 2: Data Design & Preprocessing...")
build_chapter_2(doc)

print("Building Chapter 3: Big Data System Construction...")
build_chapter_3(doc)

print("Building Chapter 4: Experiments & Results Evaluation...")
build_chapter_4(doc)

print("Building Chapter 5: Conclusion & Future Work...")
build_chapter_5(doc)

print("Building References...")
build_references(doc)

print("Building Appendices A, B, C...")
build_appendices(doc)

# ==============================================================================
# 4. LƯU BÁO CÁO WORD
# ==============================================================================
try:
    doc.save(output_docx_path)
    print("SUCCESS: BaoCao.docx has been generated successfully!")
except Exception as e:
    backup_path = r"c:\dulieulon\BaoCao_capnhat.docx"
    print(f"Warning: {output_docx_path} is currently locked by Word/WPS. Saving to {backup_path}")
    doc.save(backup_path)
    print(f"SUCCESS: {backup_path} has been generated successfully!")
