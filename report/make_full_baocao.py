# -*- coding: utf-8 -*-
"""
CHƯƠNG TRÌNH TẠO BÁO CÁO TOÀN DIỆN CHO BÀI TẬP LỚN XỬ LÝ DỮ LIỆU LỚN (CTKU122)
TRƯỜNG ĐẠI HỌC TÀI NGUYÊN VÀ MÔI TRƯỜNG HÀ NỘI
Đề tài: Hệ thống Phân tích và Giám sát Chất lượng Mạng Viễn thông (Hadoop, Spark, R)
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

template_path = r"c:\dulieulon\Mau_Bao_cao_BTL_Xu_ly_du_lieu_lon_CTKU122 (1).docx"
output_docx_path = r"c:\dulieulon\BaoCao.docx"
img_dir = r"c:\dulieulon\report"

doc = docx.Document(template_path)

# Cấu hình lề trang theo chuẩn: Trái 3.0cm, Phải 2.0cm, Trên 2.0cm, Dưới 2.0cm
for section in doc.sections:
    section.top_margin = Inches(2.0 / 2.54)
    section.bottom_margin = Inches(2.0 / 2.54)
    section.left_margin = Inches(3.0 / 2.54)
    section.right_margin = Inches(2.0 / 2.54)

# Cập nhật thông tin trang bìa
doc.paragraphs[6].text = "ĐỀ TÀI SỐ 01"
doc.paragraphs[6].runs[0].font.name = "Times New Roman"
doc.paragraphs[6].runs[0].font.size = Pt(14)
doc.paragraphs[6].runs[0].font.bold = True

doc.paragraphs[7].text = "XÂY DỰNG HỆ THỐNG PHÂN TÍCH VÀ GIÁM SÁT CHẤT LƯỢNG MẠNG VIỄN THÔNG (TELECOM QUALITY OF SERVICE ANALYTICS) TRÊN NỀN TẢNG HADOOP, SPARK VÀ R"
doc.paragraphs[7].runs[0].font.name = "Times New Roman"
doc.paragraphs[7].runs[0].font.size = Pt(16)
doc.paragraphs[7].runs[0].font.bold = True
doc.paragraphs[7].alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.paragraphs[8].text = "Lớp học phần: ĐH11.CNTT - Xử lý Dữ liệu lớn"
doc.paragraphs[8].runs[0].font.name = "Times New Roman"
doc.paragraphs[8].runs[0].font.size = Pt(13)

doc.paragraphs[9].text = "Giảng viên giảng dạy: ThS. / TS. Giảng viên phụ trách học phần"
doc.paragraphs[9].runs[0].font.name = "Times New Roman"
doc.paragraphs[9].runs[0].font.size = Pt(13)

# Cập nhật bảng sinh viên trên bìa (Table 0)
sv_list = [
    ("1", "211101001", "Nguyễn Văn A (Nhóm trưởng)", "ĐH11.CNTT"),
    ("2", "211101002", "Trần Thị B", "ĐH11.CNTT"),
    ("3", "211101003", "Lê Hoàng C", "ĐH11.CNTT"),
]
for idx, (tt, msv, hoten, lop) in enumerate(sv_list):
    r_cells = doc.tables[0].rows[idx + 1].cells
    r_cells[0].text, r_cells[1].text, r_cells[2].text, r_cells[3].text = tt, msv, hoten, lop
    for c in r_cells:
        for p in c.paragraphs:
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(11)

# Cập nhật bảng phân công nhiệm vụ (Table 5)
tasks_list = [
    ("1", "211101001", "Nguyễn Văn A", "Thiết kế kiến trúc hệ thống, cấu hình Flume Ingestion, lưu trữ HDFS và thuật toán MapReduce Java", "35%"),
    ("2", "211101002", "Trần Thị B", "Xây dựng pipeline Spark DataFrame, Spark SQL và kịch bản phân cụm dữ liệu với Apache Pig Latin", "35%"),
    ("3", "211101003", "Lê Hoàng C", "Thiết kế bảng HBase NoSQL, phân tích R & SQLite, mô hình Logistic Regression và soạn thảo báo cáo", "30%"),
]
for idx, (tt, msv, hoten, nv, tl) in enumerate(tasks_list):
    r_cells = doc.tables[5].rows[idx + 1].cells
    r_cells[0].text, r_cells[1].text, r_cells[2].text, r_cells[3].text, r_cells[4].text = tt, msv, hoten, nv, tl
    for c in r_cells:
        for p in c.paragraphs:
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(10.5)

# Cắt bỏ các đoạn outline sau Table 5 để điền nội dung toàn diện
body = doc._body._element
children = list(body)
cut_index = -1
tbl_count = 0
for i, el in enumerate(children):
    if el.tag.endswith('tbl'):
        tbl_count += 1
        if tbl_count == 6:
            cut_index = i
            break

for el in children[cut_index + 2:]:
    if el.tag.endswith('sectPr'):
        continue
    body.remove(el)

# ------------------------------------------------------------------------------
# HÀM BỔ TRỢ ĐỊNH DẠNG
# ------------------------------------------------------------------------------
def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_p(text="", bold_prefix=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=6, line_spacing=1.35, italic=False):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    if bold_prefix:
        r_b = p.add_run(bold_prefix)
        r_b.font.name = 'Times New Roman'
        r_b.font.size = Pt(13)
        r_b.font.bold = True
        r_b.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    if text:
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(13)
        r.font.italic = italic
        r.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
    return p

def add_bullet(text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.left_indent = Inches(0.25)
    if bold_prefix:
        r_b = p.add_run(bold_prefix)
        r_b.font.name = 'Times New Roman'
        r_b.font.size = Pt(13)
        r_b.font.bold = True
        r_b.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
    return p

def add_h1(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(22)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    return p

def add_h2(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
    return p

def add_h3(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.italic = True
    r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    return p

def add_tbl(headers, data, caption=None, col_widths=None):
    if caption:
        cp = doc.add_paragraph()
        cp.paragraph_format.space_before = Pt(12)
        cp.paragraph_format.space_after = Pt(4)
        cp.paragraph_format.keep_with_next = True
        c_run = cp.add_run(caption)
        c_run.font.name = 'Times New Roman'
        c_run.font.size = Pt(12)
        c_run.font.bold = True
        c_run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "E2E8F0")
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=120, right=120)
        for p in hdr_cells[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(11)
                r.font.bold = True
                r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    for r_idx, row_data in enumerate(data):
        row_cells = table.rows[r_idx + 1].cells
        bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            row_cells[c_idx].text = str(val)
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=80, bottom=80, left=100, right=100)
            for p in row_cells[c_idx].paragraphs:
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.15
                if c_idx == 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.name = 'Times New Roman'
                    r.font.size = Pt(10.5)

    if col_widths:
        for row in table.rows:
            for idx, w in enumerate(col_widths):
                row.cells[idx].width = Inches(w)

    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)
    return table

def add_fig(img_name, caption, width_in=6.0):
    img_path = os.path.join(img_dir, img_name)
    if os.path.exists(img_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run()
        run.add_picture(img_path, width=Inches(width_in))

        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_before = Pt(0)
        cp.paragraph_format.space_after = Pt(10)
        c_run = cp.add_run(caption)
        c_run.font.name = 'Times New Roman'
        c_run.font.size = Pt(11.5)
        c_run.font.italic = True
        c_run.font.bold = True
        c_run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

def add_code(code_str, caption=None):
    if caption:
        cp = doc.add_paragraph()
        cp.paragraph_format.space_before = Pt(8)
        cp.paragraph_format.space_after = Pt(2)
        cp.paragraph_format.keep_with_next = True
        c_run = cp.add_run(caption)
        c_run.font.name = 'Times New Roman'
        c_run.font.size = Pt(11)
        c_run.font.italic = True
        c_run.font.bold = True

    tbl = doc.add_table(rows=1, cols=1)
    tbl.style = 'Table Grid'
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F1F5F9")
    set_cell_margins(cell, top=100, bottom=100, left=150, right=150)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(code_str)
    r.font.name = 'Consolas'
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)

def add_callout(title, text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.style = 'Table Grid'
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "EFF6FF")
    set_cell_margins(cell, top=120, bottom=120, left=160, right=160)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.25
    r_t = p.add_run(f"📌 {title}: ")
    r_t.font.name = 'Times New Roman'
    r_t.font.size = Pt(11.5)
    r_t.font.bold = True
    r_t.font.color.rgb = RGBColor(0x1D, 0x4E, 0xD8)

    r_txt = p.add_run(text)
    r_txt.font.name = 'Times New Roman'
    r_txt.font.size = Pt(11)
    r_txt.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)

print("Base setup completed.")
