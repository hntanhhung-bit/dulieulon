# -*- coding: utf-8 -*-
"""
MODULE 1: BỔ TRỢ ĐỊNH DẠNG, MỤC LỤC, MỞ ĐẦU VÀ CHƯƠNG 1
Tuân thủ quy chuẩn học phần Xử lý dữ liệu lớn (CTKU122 - HUNRE)
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

img_dir = r"c:\dulieulon\report"

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_p(doc, text="", bold_prefix=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=6, line_spacing=1.35, italic=False):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    if bold_prefix:
        r_b = p.add_run(bold_prefix)
        r_b.font.name = 'Times New Roman'
        r_b.font.size = Pt(13)
        r_b.font.bold = True
        r_b.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    if text:
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(13)
        r.font.italic = italic
        r.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
    return p

def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.left_indent = Inches(0.25)
    if bold_prefix:
        r_b = p.add_run(bold_prefix)
        r_b.font.name = 'Times New Roman'
        r_b.font.size = Pt(13)
        r_b.font.bold = True
        r_b.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
    return p

def add_h1(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(24)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.keep_with_next = True
    pPr = p._p.get_or_add_pPr()
    pPr.append(parse_xml(f'<w:outlineLvl {nsdecls("w")} w:val="0"/>'))
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    return p

def add_h2(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    pPr = p._p.get_or_add_pPr()
    pPr.append(parse_xml(f'<w:outlineLvl {nsdecls("w")} w:val="1"/>'))
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
    return p

def add_h3(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    pPr = p._p.get_or_add_pPr()
    pPr.append(parse_xml(f'<w:outlineLvl {nsdecls("w")} w:val="2"/>'))
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.italic = True
    r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    return p

def add_tbl(doc, headers, data, caption=None, col_widths=None):
    if caption:
        cp = doc.add_paragraph()
        cp.paragraph_format.space_before = Pt(12)
        cp.paragraph_format.space_after = Pt(4)
        cp.paragraph_format.keep_with_next = True
        c_run = cp.add_run(caption)
        c_run.font.name = 'Times New Roman'
        c_run.font.size = Pt(12)
        c_run.font.bold = True
        c_run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "E2E8F0")
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=120, right=120)
        for p in hdr_cells[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(11)
                r.font.bold = True
                r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    for r_idx, row_data in enumerate(data):
        row_cells = table.rows[r_idx + 1].cells
        bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            row_cells[c_idx].text = str(val)
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=80, bottom=80, left=100, right=100)
            for p in row_cells[c_idx].paragraphs:
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.15
                if c_idx == 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.name = 'Times New Roman'
                    r.font.size = Pt(10.5)

    if col_widths:
        for row in table.rows:
            for idx, w in enumerate(col_widths):
                row.cells[idx].width = Inches(w)

    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)
    return table

def add_fig(doc, img_name, caption, width_in=5.8):
    img_path = os.path.join(img_dir, img_name)
    if os.path.exists(img_path):
        actual_width = min(float(width_in), 5.85)

        # Đoạn văn chứa hình ảnh: SINGLE LINE SPACING để ngăn chữ đè ảnh 100%
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        p.paragraph_format.space_before = Pt(20)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run()
        run.add_picture(img_path, width=Inches(actual_width))

        # Đoạn văn chú thích tên hình
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.line_spacing = 1.15
        cp.paragraph_format.space_before = Pt(4)
        cp.paragraph_format.space_after = Pt(18)
        c_run = cp.add_run(caption)
        c_run.font.name = 'Times New Roman'
        c_run.font.size = Pt(11.5)
        c_run.font.italic = True
        c_run.font.bold = True
        c_run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

def add_code(doc, code_str, caption=None):
    if caption:
        cp = doc.add_paragraph()
        cp.paragraph_format.space_before = Pt(8)
        cp.paragraph_format.space_after = Pt(2)
        cp.paragraph_format.keep_with_next = True
        c_run = cp.add_run(caption)
        c_run.font.name = 'Times New Roman'
        c_run.font.size = Pt(11)
        c_run.font.italic = True
        c_run.font.bold = True

    tbl = doc.add_table(rows=1, cols=1)
    tbl.style = 'Table Grid'
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F1F5F9")
    set_cell_margins(cell, top=100, bottom=100, left=150, right=150)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(code_str)
    r.font.name = 'Consolas'
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)

def add_callout(doc, title, text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.style = 'Table Grid'
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "EFF6FF")
    set_cell_margins(cell, top=120, bottom=120, left=160, right=160)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.25
    r_t = p.add_run(f"📌 {title}: ")
    r_t.font.name = 'Times New Roman'
    r_t.font.size = Pt(11.5)
    r_t.font.bold = True
    r_t.font.color.rgb = RGBColor(0x1D, 0x4E, 0xD8)

    r_txt = p.add_run(text)
    r_txt.font.name = 'Times New Roman'
    r_txt.font.size = Pt(11)
    r_txt.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)

# ==============================================================================
# PHẦN MỤC LỤC & DANH MỤC
# ==============================================================================
def build_toc_and_lists(doc):
    doc.add_page_break()
    add_h1(doc, "MỤC LỤC")

    toc_items = [
        ("LỜI NÓI ĐẦU", "5"),
        ("CHƯƠNG 1. TỔNG QUAN VÀ BÀI TOÁN XỬ LÝ DỮ LIỆU LỚN", "7"),
        ("  1.1. Bối cảnh và lý do chọn đề tài", "7"),
        ("  1.2. Phân tích 5 đặc trưng Big Data (5V) trong bài toán viễn thông", "8"),
        ("  1.3. Khảo sát các công nghệ Big Data áp dụng trong đề tài", "10"),
        ("  1.4. Thiết kế kiến trúc tổng thể hệ thống (Kiến trúc Lambda)", "12"),
        ("CHƯƠNG 2. THIẾT KẾ DỮ LIỆU VÀ TIỀN XỬ LÝ", "14"),
        ("  2.1. Nguồn dữ liệu và phân loại dữ liệu viễn thông", "14"),
        ("  2.2. Phân định ranh giới xử lý: Gần thời gian thực vs Phân tích tĩnh", "15"),
        ("  2.3. Khảo sát cấu trúc các tập dữ liệu thực nghiệm", "16"),
        ("  2.4. Đánh giá chất lượng dữ liệu và các vấn đề dị thường", "18"),
        ("  2.5. Quy trình làm sạch dữ liệu và kỹ thuật tạo đặc trưng", "19"),
        ("CHƯƠNG 3. XÂY DỰNG HỆ THỐNG XỬ LÝ DỮ LIỆU LỚN", "21"),
        ("  3.1. Thiết lập môi trường triển khai thực nghiệm", "21"),
        ("  3.2. Tổ chức lưu trữ phân tán và cơ chế dung lỗi trên HDFS", "22"),
        ("  3.3. Thiết kế pipeline luồng dữ liệu đầu cuối (End-to-End)", "25"),
        ("  3.4. Cài đặt pipeline phân tích bằng Apache Spark & Spark SQL", "27"),
        ("  3.5. Các thành phần và thuật toán chuyên biệt trong đề tài", "29"),
        ("    3.5.1. Thuật toán MapReduce tối ưu tính KPI viễn thông", "29"),
        ("    3.5.2. Kịch bản ETL và phân cụm trạm lỗi với Apache Pig Latin", "31"),
        ("    3.5.3. Thiết kế bảng Apache HBase phục vụ truy cập ngẫu nhiên", "33"),
        ("    3.5.4. Mô phỏng phân tích nghiệp vụ tương tác với BigSheets", "35"),
        ("    3.5.5. Mô hình Hồi quy Logistic dự báo khiếu nại bằng R & SQLite", "37"),
        ("  3.6. Cơ chế lưu trữ và xuất kết quả đa tầng", "39"),
        ("CHƯƠNG 4. THỰC NGHIỆM VÀ ĐÁNH GIÁ KẾT QUẢ", "40"),
        ("  4.1. Kịch bản và dữ liệu thực nghiệm", "40"),
        ("  4.2. Kết quả phân tích chất lượng mạng viễn thông", "41"),
        ("  4.3. Đánh giá hiệu năng và so sánh công nghệ", "46"),
        ("  4.4. Thảo luận kết quả và 5 Khuyến nghị hành động tối ưu mạng NOC", "48"),
        ("CHƯƠNG 5. KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN", "52"),
        ("  5.1. Kết luận các kết quả đạt được", "52"),
        ("  5.2. Hạn chế của đề tài", "53"),
        ("  5.3. Hướng nghiên cứu và phát triển tiếp theo", "53"),
        ("TÀI LIỆU THAM KHẢO", "55"),
        ("PHỤ LỤC", "56"),
        ("  Phụ lục A: Hướng dẫn cài đặt và vận hành hệ thống", "56"),
        ("  Phụ lục B: Trích đoạn mã nguồn cốt lõi", "58"),
        ("  Phụ lục C: Nhật ký thực thi và ảnh chụp màn hình kiểm chứng", "62"),
    ]

    for title, pg in toc_items:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        r_t = p.add_run(title)
        r_t.font.name = "Times New Roman"
        r_t.font.size = Pt(12)
        if title.startswith("CHƯƠNG") or title in ["LỜI NÓI ĐẦU", "TÀI LIỆU THAM KHẢO", "PHỤ LỤC"]:
            r_t.font.bold = True
        dots_len = max(3, 85 - len(title))
        r_dots = p.add_run(" " + "." * dots_len + " ")
        r_dots.font.name = "Times New Roman"
        r_dots.font.size = Pt(11)
        r_dots.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)
        r_pg = p.add_run(pg)
        r_pg.font.name = "Times New Roman"
        r_pg.font.size = Pt(12)
        r_pg.font.bold = True

    doc.add_page_break()
    add_h1(doc, "DANH MỤC HÌNH VẼ")
    figures_list = [
        ("Hình 1.1", "Sơ đồ kiến trúc tổng thể hệ thống (Lambda Architecture)", "13"),
        ("Hình 2.1", "Sơ đồ thực thể quan hệ dữ liệu viễn thông (ERD Schema)", "17"),
        ("Hình 3.1", "Pipeline luồng dữ liệu đầu cuối từ Ingestion đến Analytics", "26"),
        ("Hình 3.2", "Sơ đồ luồng thuật toán MapReduce với Composite Key và Custom Partitioner", "30"),
        ("Hình 3.3", "Giao diện mô phỏng Dashboard phân tích nghiệp vụ BigSheets", "36"),
        ("Hình 4.1", "Biểu đồ hộp phân bố độ trễ theo thế hệ công nghệ mạng (3G, 4G, 5G)", "42"),
        ("Hình 4.2", "Biểu đồ tương quan giữa Tỷ lệ mất gói (%) và Tỷ lệ rớt cuộc gọi (%)", "43"),
        ("Hình 4.3", "Biểu đồ tỷ lệ rớt cuộc gọi theo Tỉnh thành và Công nghệ mạng", "44"),
        ("Hình 4.4", "Biểu đồ bong bóng đa chiều tương quan Độ trễ và Mất gói theo Trạm BTS", "45"),
        ("Hình 4.5", "Biểu đồ xếp hạng cảnh báo Top 10 trạm BTS chất lượng kém nhất cần bảo trì", "45"),
        ("Hình 4.6", "Ảnh chụp màn hình môi trường thực thi và phân tích mô hình học máy trên RStudio", "46"),
    ]
    for h_id, h_name, h_pg in figures_list:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        r_id = p.add_run(f"{h_id}: ")
        r_id.font.name = "Times New Roman"
        r_id.font.size = Pt(12)
        r_id.font.bold = True
        r_nm = p.add_run(h_name)
        r_nm.font.name = "Times New Roman"
        r_nm.font.size = Pt(12)
        dots_len = max(3, 80 - len(h_id) - len(h_name))
        r_dots = p.add_run(" " + "." * dots_len + " ")
        r_dots.font.name = "Times New Roman"
        r_dots.font.size = Pt(11)
        r_dots.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)
        r_pg = p.add_run(h_pg)
        r_pg.font.name = "Times New Roman"
        r_pg.font.size = Pt(12)
        r_pg.font.bold = True

    add_h1(doc, "DANH MỤC BẢNG BIỂU")
    tables_list = [
        ("Bảng 2.1", "Khảo sát cấu trúc 4 tập dữ liệu viễn thông thực nghiệm", "16"),
        ("Bảng 2.2", "Ma trận rủi ro và phương án xử lý chất lượng dữ liệu đo kiểm", "18"),
        ("Bảng 3.1", "Thiết kế bảng Apache HBase lưu trữ chỉ số KPI viễn thông độ trễ thấp", "34"),
        ("Bảng 4.1", "Bảng tổng hợp KPI Top 10 trạm phát sóng suy giảm chất lượng nghiêm trọng", "41"),
        ("Bảng 4.2", "Bảng đối chiếu định lượng hiệu năng và tài nguyên giữa Pig, MapReduce, Spark và R", "47"),
        ("Bảng 4.3", "Ma trận 5 khuyến nghị hành động kỹ thuật tối ưu hóa mạng viễn thông", "49"),
    ]
    for b_id, b_name, b_pg in tables_list:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        r_id = p.add_run(f"{b_id}: ")
        r_id.font.name = "Times New Roman"
        r_id.font.size = Pt(12)
        r_id.font.bold = True
        r_nm = p.add_run(b_name)
        r_nm.font.name = "Times New Roman"
        r_nm.font.size = Pt(12)
        dots_len = max(3, 80 - len(b_id) - len(b_name))
        r_dots = p.add_run(" " + "." * dots_len + " ")
        r_dots.font.name = "Times New Roman"
        r_dots.font.size = Pt(11)
        r_dots.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)
        r_pg = p.add_run(b_pg)
        r_pg.font.name = "Times New Roman"
        r_pg.font.size = Pt(12)
        r_pg.font.bold = True

    add_h1(doc, "DANH MỤC TỪ VIẾT TẮT")
    abbrs = [
        ("BTS", "Base Transceiver Station - Trạm thu phát sóng di động mặt đất"),
        ("CDR / xDR", "Call Detail Record / Data Detail Record - Bản ghi chi tiết cuộc gọi và dữ liệu"),
        ("KPI", "Key Performance Indicator - Chỉ số đo lường hiệu quả hoạt động cốt lõi"),
        ("QoS / QoE", "Quality of Service / Quality of Experience - Chất lượng dịch vụ và trải nghiệm người dùng"),
        ("SLA", "Service Level Agreement - Cam kết chất lượng dịch vụ"),
        ("HDFS", "Hadoop Distributed File System - Hệ thống tệp phân tán Hadoop"),
        ("YARN", "Yet Another Resource Negotiator - Hệ thống quản lý tài nguyên phân tán Hadoop"),
        ("RDD", "Resilient Distributed Dataset - Tập dữ liệu phân tán có khả năng phục hồi"),
        ("NoSQL", "Not Only SQL - Cơ sở dữ liệu phi quan hệ phân tán"),
        ("SON", "Self-Organizing Networks - Mạng tự tối ưu hóa vô tuyến"),
        ("NOC / SOC", "Network Operations Center / Service Operations Center - Trung tâm điều hành mạng và dịch vụ"),
        ("MEC", "Multi-access Edge Computing - Tính toán biên đa truy cập"),
        ("AQM / WRED", "Active Queue Management / Weighted Random Early Detection - Kiểm soát hàng đợi chủ động"),
    ]
    for ab, desc in abbrs:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        r_a = p.add_run(f"{ab:<12} : ")
        r_a.font.name = "Times New Roman"
        r_a.font.size = Pt(12)
        r_a.font.bold = True
        r_d = p.add_run(desc)
        r_d.font.name = "Times New Roman"
        r_d.font.size = Pt(12)

# ==============================================================================
# LỜI NÓI ĐẦU / MỞ ĐẦU
# ==============================================================================
def build_introduction(doc):
    doc.add_page_break()
    add_h1(doc, "LỜI NÓI ĐẦU")
    add_p(doc, "Trong kỷ nguyên bùng nổ của cuộc Cách mạng Công nghiệp lần thứ tư, viễn thông di động đóng vai trò là mạch máu thông tin huyết mạch của toàn bộ nền kinh tế số. Sự phổ cập vũ bão của mạng di động thế hệ thứ tư (4G LTE), thế hệ thứ năm (5G NR) cùng với sự gia tăng đột biến của các thiết bị Internet vạn vật (IoT) và các dịch vụ truyền tải trực tuyến độ phân giải cao đã tạo ra áp lực lưu lượng chưa từng có đối với cơ sở hạ tầng mạng của các nhà khai thác viễn thông.")
    add_p(doc, "Một mạng viễn thông quy mô quốc gia thông thường sở hữu hàng chục nghìn trạm phát sóng di động (Base Transceiver Station - BTS). Mỗi trạm BTS liên tục ghi nhận và phát sinh các dòng dữ liệu đo kiểm telemetry, bản ghi chi tiết cuộc gọi (Call Detail Record - CDR), bản ghi phiên dữ liệu (Data Detail Record - xDR) và nhật ký hoạt động hệ điều hành với tần suất hàng trăm nghìn sự kiện mỗi giây. Khối lượng dữ liệu này nhanh chóng tích lũy lên tới hàng chục Terabyte mỗi ngày, mang đầy đủ các đặc trưng kinh điển của Dữ liệu lớn (Big Data): Khối lượng khổng lồ (Volume), Tốc độ sinh và truyền tức thời (Velocity), Đa dạng định dạng (Variety), Độ nhiễu và bất định cao (Veracity), cùng tiềm năng khai thác Giá trị to lớn (Value).")
    add_p(doc, "Trước áp lực cạnh tranh khốc liệt trên thị trường, chất lượng dịch vụ (Quality of Service - QoS) và chất lượng trải nghiệm của khách hàng (Quality of Experience - QoE) đã trở thành yếu tố sống còn quyết định sự thành bại của doanh nghiệp viễn thông. Các hệ thống cơ sở dữ liệu quan hệ truyền thống (RDBMS) hoàn toàn bất lực và sụp đổ trước khối lượng dữ liệu khổng lồ này, không thể cung cấp khả năng phân tích đa chiều hoặc cảnh báo sự cố tức thời.")
    add_p(doc, "Xuất phát từ thực tiễn cấp bách đó, nhóm nghiên cứu đã lựa chọn đề tài: \"Xây dựng hệ thống phân tích và giám sát chất lượng mạng viễn thông (Telecom Quality of Service Analytics) trên nền tảng Hadoop, Spark và R\". Đề tài hướng tới việc giải quyết triệt để 3 mục tiêu kỹ thuật cốt lõi:", bold_prefix="Mục tiêu cốt lõi của đề tài: ")
    add_bullet(doc, "Giảm thiểu độ trễ mạng (Latency): Phát hiện chính xác các trạm phát sóng bị nghẽn đường truyền truyền dẫn Backhaul và quá tải hàng đợi vô tuyến, hỗ trợ định tuyến tối ưu lưu lượng.", bold_prefix="1. ")
    add_bullet(doc, "Giảm tỷ lệ mất gói (Packet Loss): Cô lập các cell vô tuyến bị can nhiễu đồng kênh, lỗi tham số công suất phát hoặc tràn bộ đệm thiết bị định tuyến.", bold_prefix="2. ")
    add_bullet(doc, "Giảm tỷ lệ rớt cuộc gọi (Call Drop Rate - CDR): Tối ưu hóa vùng phủ sóng, hỗ trợ tinh chỉnh tham số biên chuyển giao (Handover Margin) và góc ngửa anten, kéo giảm tỷ lệ rớt cuộc gọi về dưới ngưỡng cam kết chất lượng dịch vụ SLA (2.0%).", bold_prefix="3. ")

    add_p(doc, "Để đạt được các mục tiêu trên, đề tài đã thiết kế và triển khai một hệ sinh thái Big Data tích hợp hoàn chỉnh theo kiến trúc Lambda hiện đại, bao gồm: Apache Flume (thu thập luồng log tự động), Hadoop HDFS (kho lưu trữ Data Lake dung lỗi), Apache Spark & Spark SQL (xử lý dữ liệu phân tán tốc độ cao trong bộ nhớ), Hadoop MapReduce Java (thuật toán tối ưu hóa với Composite Key và Custom Partitioner), Apache Pig Latin (kịch bản ETL và phân cụm trạm lỗi), Apache HBase (cơ sở dữ liệu NoSQL truy cập ngẫu nhiên độ trễ thấp < 10ms), BigSheets (mô phỏng bảng tính phân tích nghiệp vụ tương tác No-Code) và Ngôn ngữ R kết hợp SQLite (mô hình học máy Logistic Regression dự báo nguy cơ phát sinh khiếu nại khách hàng).")

    add_p(doc, "Bố cục của báo cáo Bài tập lớn được tổ chức khoa học thành 5 chương chính như sau:", bold_prefix="Cấu trúc báo cáo: ")
    add_bullet(doc, "Trình bày bối cảnh thực tiễn, phân tích 5 đặc trưng Big Data (5V), khảo sát các công nghệ áp dụng và thiết kế kiến trúc hệ thống tổng thể theo mô hình Lambda Architecture.", bold_prefix="Chương 1 - Tổng quan và bài toán xử lý dữ liệu lớn: ")
    add_bullet(doc, "Phân loại dữ liệu viễn thông, phân định ranh giới xử lý thời gian thực vs xử lý lô, khảo sát cấu trúc 4 tập dữ liệu thực nghiệm, đánh giá ma trận rủi ro chất lượng dữ liệu và quy trình làm sạch, kỹ thuật tạo đặc trưng.", bold_prefix="Chương 2 - Thiết kế dữ liệu và tiền xử lý: ")
    add_bullet(doc, "Mô tả chi tiết môi trường thực nghiệm, thiết kế lưu trữ HDFS, tính toán dung lượng sao chép RF=3, xây dựng pipeline Flume, cài đặt Spark DataFrame, thuật toán MapReduce tối ưu, kịch bản Pig Latin, thiết kế HBase NoSQL, quy trình BigSheets và mô hình Hồi quy Logistic bằng R.", bold_prefix="Chương 3 - Xây dựng hệ thống xử lý dữ liệu lớn: ")
    add_bullet(doc, "Công bố và thảo luận kết quả thực nghiệm từ 12,000 phiên đo kiểm, phân tích các biểu đồ chuyên sâu (Hình 4.1 đến Hình 4.6), đối chiếu định lượng hiệu năng giữa 4 công nghệ (Bảng 4.2) và đề xuất 5 khuyến nghị hành động chiến lược cho Trung tâm điều hành mạng NOC.", bold_prefix="Chương 4 - Thực nghiệm và đánh giá kết quả: ")
    add_bullet(doc, "Tổng kết các kết quả đạt được, chỉ ra các hạn chế tồn đọng và định hướng mở rộng nghiên cứu trong tương lai (Streaming với Kafka & Flink, ứng dụng AIOps).", bold_prefix="Chương 5 - Kết luận và hướng phát triển: ")

# ==============================================================================
# CHƯƠNG 1: TỔNG QUAN VÀ BÀI TOÁN XỬ LÝ DỮ LIỆU LỚN
# ==============================================================================
def build_chapter_1(doc):
    doc.add_page_break()
    add_h1(doc, "CHƯƠNG 1. TỔNG QUAN VÀ BÀI TOÁN XỬ LÝ DỮ LIỆU LỚN")
    
    add_h2(doc, "1.1. Bối cảnh và lý do chọn đề tài")
    add_p(doc, "Trong thời đại số hóa toàn diện, ngành công nghiệp viễn thông đang đóng vai trò là hạ tầng kết nối nền tảng cho sự phát triển của xã hội số và kinh tế số. Sự chuyển dịch mạnh mẽ từ các dịch vụ thoại/SMS truyền thống sang các dịch vụ dữ liệu băng rộng tốc độ cao (Data Services) như video streaming độ phân giải cao 4K/8K, điện toán đám mây, hội nghị truyền hình trực tuyến, game thực tế ảo và xe tự hành đã tạo nên một làn sóng bùng nổ lưu lượng truy cập trên các mạng vô tuyến di động.")
    add_p(doc, "Một doanh nghiệp viễn thông tiêu biểu tại Việt Nam quản lý từ 30,000 đến hơn 60,000 trạm BTS phân bố trải dài trên khắp 63 tỉnh thành cả nước. Mỗi trạm phát sóng bao gồm từ 3 đến 6 cell vô tuyến phủ sóng ở các dải tần số khác nhau (900 MHz, 1800 MHz, 2100 MHz, 2600 MHz) phục vụ đồng thời hàng triệu thuê bao di động. Tại mỗi cell, các cảm biến phần cứng và phần mềm nhúng liên tục đo kiểm và ghi nhận hàng trăm tham số chất lượng dịch vụ (Quality of Service - QoS) theo chu kỳ từng giây.")
    add_p(doc, "Thách thức lớn nhất mà các kỹ sư vận hành mạng tại Trung tâm Điều hành Mạng (Network Operations Center - NOC) đang phải đối mặt bao gồm:")
    add_bullet(doc, "Khối lượng dữ liệu khổng lồ sinh ra vượt quá năng lực xử lý của các giải pháp lưu trữ RDBMS truyền thống, dẫn đến việc dữ liệu bị xóa bỏ hoặc nén mù mờ mà không kịp trích xuất giá trị thông tin.")
    add_bullet(doc, "Hiện tượng nghẽn mạng cục bộ và suy giảm chất lượng dịch vụ diễn ra âm thầm, chỉ được phát hiện khi khách hàng đã gặp sự cố rớt cuộc gọi nhiều lần và gọi điện phàn nàn lên tổng đài Chăm sóc khách hàng (CSKH), gây tổn hại nghiêm trọng đến uy tín thương hiệu.")
    add_bullet(doc, "Chi phí đầu tư nâng cấp hạ tầng (CAPEX) rất lớn nhưng thiếu căn cứ định lượng chính xác: Nhà mạng không thể biết rõ trạm BTS bị suy giảm chất lượng là do thiếu tần số vô tuyến, do nghẽn đường truyền cáp quang Backhaul hay do cấu hình sai góc ngửa anten vô tuyến.")
    add_p(doc, "Do đó, việc nghiên cứu và xây dựng một hệ thống phân tích, giám sát chất lượng mạng tự động trên nền tảng các công nghệ Xử lý Dữ liệu lớn (Big Data Technologies) là một yêu cầu cấp thiết, có giá trị ứng dụng thực tiễn cao đối với các doanh nghiệp viễn thông hiện nay.")

    add_h2(doc, "1.2. Phân tích 5 đặc trưng Big Data (5V) trong bài toán viễn thông")
    add_p(doc, "Bài toán phân tích dữ liệu mạng viễn thông thể hiện đầy đủ và sống động 5 đặc trưng kinh điển của dữ liệu lớn (5V Model):")
    
    add_bullet(doc, "Volume phản ánh quy mô dữ liệu khổng lồ. Với một mạng lưới quản lý khoảng 50,000 trạm BTS, mỗi trạm có trung bình 4 cell vô tuyến đo kiểm liên tục theo chu kỳ 5 giây/lần, toàn mạng sinh ra xấp xỉ 3.5 tỷ sự kiện đo kiểm mỗi ngày. Kèm theo đó là hàng trăm triệu bản ghi chi tiết cuộc gọi (CDR) và phiên truy cập Internet (xDR). Khối lượng dữ liệu thô sinh ra dao động từ 10 đến 25 Terabyte mỗi ngày, tích lũy qua các năm lên tới quy mô hàng chục Petabyte.", bold_prefix="1. Volume (Khối lượng dữ liệu khổng lồ): ")
    
    add_bullet(doc, "Velocity phản ánh tốc độ sinh ra, truyền tải và yêu cầu xử lý dữ liệu cực cao. Dữ liệu từ các trạm BTS được truyền liên tục về máy chủ trung tâm qua các luồng stream thời gian thực với thông lượng vượt mức 500,000 sự kiện/giây vào các khung giờ cao điểm (19h00 - 22h00). Hệ thống giám sát NOC đòi hỏi khả năng tiếp nhận, xử lý và kích hoạt cảnh báo trạm mất sóng hoặc rớt cuộc gọi đột biến trong vòng dưới 1 giây để các kỹ sư kịp thời can thiệp.", bold_prefix="2. Velocity (Tốc độ sinh và truyền dữ liệu tức thời): ")

    add_bullet(doc, "Variety thể hiện ở sự đa dạng phức tạp về nguồn gốc và định dạng dữ liệu. Hệ sinh thái viễn thông là sự pha trộn của nhiều công nghệ mạng (2G/3G kế thừa, 4G LTE băng rộng, 5G độ trễ cực thấp) đến từ nhiều nhà cung cấp phần cứng khác nhau (Ericsson, Huawei, Nokia, ZTE). Định dạng dữ liệu bao gồm: dữ liệu có cấu trúc (bảng danh mục trạm BTS, danh mục cell, tọa độ GPS), dữ liệu bán cấu trúc (file nhật ký Syslog chuẩn RFC-5424, file cấu hình tham số vô tuyến định dạng XML/JSON), và dữ liệu luồng phi cấu trúc (luồng gói tin NetFlow/IPFIX, telemetry đo kiểm lớp vật lý).", bold_prefix="3. Variety (Đa dạng về chủng loại và định dạng dữ liệu): ")

    add_bullet(doc, "Veracity thể hiện tính xác thực, độ tin cậy và mức độ nhiễu của dữ liệu. Do tín hiệu vô tuyến truyền lan trong môi trường không gian tự do chịu ảnh hưởng nặng nề bởi thời tiết mưa bão, che chắn địa hình đô thị, can nhiễu đồng kênh và suy hao đa đường (multipath fading), dữ liệu đo kiểm viễn thông thường xuyên xuất hiện các giá trị ngoại lai dị thường (outliers). Bên cạnh đó, hiện tượng lệch xung đồng hồ NTP giữa trạm BTS và máy chủ trung tâm có thể dẫn đến các bản ghi đo kiểm có độ trễ âm hoặc thông lượng âm vô lý. Khâu làm sạch và xác thực dữ liệu là điều kiện tiên quyết trước khi đưa vào tính toán KPI.", bold_prefix="4. Veracity (Tính xác thực và độ nhiễu của dữ liệu): ")

    add_bullet(doc, "Value là giá trị kinh tế và vận hành to lớn mà hệ thống mang lại. Khai thác hiệu quả nguồn dữ liệu này giúp nhà mạng: (1) Nâng cao chất lượng trải nghiệm mạng QoE cho người dùng; (2) Kéo giảm tỷ lệ rời mạng của thuê bao (Churn Reduction); (3) Chủ động chăm sóc khách hàng trước khi phát sinh khiếu nại gay gắt; (4) Tối ưu hóa chi phí đầu tư vốn (CAPEX/OPEX) bằng cách chỉ nâng cấp đúng các trạm bị nghẽn thực sự.", bold_prefix="5. Value (Giá trị nghiệp vụ và kinh tế cốt lõi): ")

    add_h2(doc, "1.3. Khảo sát các công nghệ Big Data áp dụng trong đề tài")
    add_p(doc, "Để giải quyết toàn diện bài toán từ khâu thu thập, lưu trữ, xử lý hàng loạt đến truy cập ngẫu nhiên và phân tích dự báo, đề tài đã nghiên cứu và tích hợp các công nghệ Big Data hàng đầu:")
    add_bullet(doc, "Apache Hadoop & HDFS: Nền tảng cốt lõi cung cấp kho dữ liệu phân tán (Data Lake) với khả năng mở rộng quy mô tuyến tính trên các máy chủ phần cứng tiêu chuẩn (Commodity Hardware). HDFS phân chia tệp dữ liệu thành các khối (Block) 128 MB và tự động nhân bản 3 lần (Replication Factor = 3) trên các node khác nhau, bảo đảm an toàn dữ liệu tuyệt đối khi xảy ra sự cố phần cứng.", bold_prefix="• Apache Hadoop & HDFS: ")
    add_bullet(doc, "Apache Flume: Dịch vụ thu thập và tổng hợp dữ liệu luồng phân tán mạnh mẽ, chuyên dụng cho việc nạp các tệp log và sự kiện từ hàng nghìn trạm BTS vào HDFS và HBase với cơ chế giao dịch hai pha (Two-Phase Commit) chống mất mát dữ liệu.", bold_prefix="• Apache Flume: ")
    add_bullet(doc, "Apache Spark & Spark SQL: Động cơ tính toán phân tán thế hệ mới xử lý dữ liệu hoàn toàn trên bộ nhớ RAM (In-Memory Processing). Spark sở hữu bộ tối ưu hóa truy vấn Catalyst Optimizer và động cơ sinh mã nhị phân Tungsten Engine, cho phép thực thi các phép biến đổi dữ liệu, Join đa bảng và tính toán thống kê nhanh gấp 10 đến 50 lần so với mô hình MapReduce truyền thống.", bold_prefix="• Apache Spark & Spark SQL: ")
    add_bullet(doc, "Hadoop MapReduce (Java): Mô hình lập trình phân tán chuẩn mực của Hadoop, tối ưu hóa cho các tác vụ xử lý lô khổng lồ chạy qua đêm (Batch Jobs). Nhóm đã thiết kế thuật toán MapReduce chuyên biệt với Khóa phức hợp (Composite Key) và Bộ phân vùng tùy biến (Custom Partitioner) để xử lý KPI viễn thông.", bold_prefix="• Hadoop MapReduce: ")
    add_bullet(doc, "Apache Pig & Pig Latin: Nền tảng xử lý dữ liệu luồng cấp cao với ngôn ngữ khai báo Pig Latin, giúp rút ngắn thời gian xây dựng các kịch bản trích xuất, chuyển đổi và nạp dữ liệu (ETL) phức tạp mà không cần viết hàng trăm dòng mã Java.", bold_prefix="• Apache Pig: ")
    add_bullet(doc, "Apache HBase: Cơ sở dữ liệu NoSQL phân tán hướng cột (Column-oriented), vận hành trực tiếp trên nền tảng HDFS, cung cấp khả năng đọc/ghi ngẫu nhiên hàng triệu bản ghi theo thời gian thực với độ trễ cực thấp (< 10 mili-giây).", bold_prefix="• Apache HBase: ")
    add_bullet(doc, "IBM BigSheets Simulator: Giải pháp mô phỏng giao diện bảng tính trực quan (Spreadsheet UI), cho phép các chuyên viên phân tích nghiệp vụ viễn thông (Business Analysts) thao tác khám phá, lọc dữ liệu và xây dựng báo cáo điều hành No-Code.", bold_prefix="• IBM BigSheets Simulator: ")
    add_bullet(doc, "Ngôn ngữ R & SQLite: Môi trường phân tích thống kê và Khoa học dữ liệu chuyên sâu, phục vụ xây dựng mô hình Học máy Hồi quy Logistic (Logistic Regression) để dự báo nguy cơ phát sinh khiếu nại khách hàng và trực quan hóa chuyên sâu bằng thư viện ggplot2.", bold_prefix="• Ngôn ngữ R & SQLite: ")

    add_h2(doc, "1.4. Thiết kế kiến trúc tổng thể hệ thống (Kiến trúc Lambda)")
    add_p(doc, "Hệ thống được thiết kế dựa trên mô hình Kiến trúc Lambda (Lambda Architecture) chuẩn mực trong kỹ nghệ dữ liệu lớn. Kiến trúc kết hợp hài hòa giữa hai luồng xử lý: luồng xử lý lô quy mô lớn (Batch Layer) đảm bảo tính chính xác toàn vẹn và luồng xử lý tốc độ cao gần thời gian thực (Speed Layer) đảm bảo tính kịp thời cho các cảnh báo vận hành.")

    add_fig(
        doc,
        img_name="hinh_1_1_kien_truc_tong_the.png",
        caption="Hình 1.1: Sơ đồ kiến trúc tổng thể hệ thống (Lambda Architecture)",
        width_in=6.5
    )

    add_p(doc, "Phân tích và nhận xét chi tiết kiến trúc hệ thống thể hiện trong Hình 1.1:", bold_prefix="Nhận xét và phân tích kiến trúc Hình 1.1: ")
    add_p(doc, "Kiến trúc hệ thống bao gồm 5 phân tầng chức năng liên hoàn:")
    add_bullet(doc, "Tầng Nguồn dữ liệu (Data Sources Layer): Tiếp nhận luồng dữ liệu thô đa dạng từ hơn 50,000 trạm BTS vô tuyến 3G/4G/5G, các bản ghi cước phiên CDR/xDR và dữ liệu phản ánh khiếu nại từ hệ thống Chăm sóc khách hàng CRM.")
    add_bullet(doc, "Tầng Thu thập & Lưu trữ (Ingestion & Storage Layer): Apache Flume đóng vai trò là cổng thu nhận trung tâm, sử dụng FileChannel bền vững để nạp dữ liệu an toàn vào kho lưu trữ phân tán Hadoop HDFS (Data Lake) với kích thước khối 128 MB và hệ số nhân bản 3 lần.")
    add_bullet(doc, "Tầng Tính toán (Compute Layer): Chia làm hai nhánh xử lý song song. Nhánh Batch Layer vận hành các tác vụ MapReduce Java, Pig Latin và Spark Batch chạy định kỳ theo giờ/ngày để tính toán các chỉ số KPI dài hạn. Nhánh Speed Layer sử dụng Spark Streaming tiếp nhận luồng sự kiện nóng để phát hiện cell nghẽn tức thời.")
    add_bullet(doc, "Tầng Phục vụ (Serving Layer): Apache HBase NoSQL lưu trữ các chỉ số KPI thời gian thực phục vụ truy vấn ngẫu nhiên độ trễ thấp (< 10 ms). Cơ sở dữ liệu SQLite lưu trữ dữ liệu tổng hợp phục vụ các mô hình phân tích R.")
    add_bullet(doc, "Tầng Báo cáo & Phân tích (Analytics & BI Layer): Cung cấp giao diện BigSheets Dashboard tương tác cho người dùng nghiệp vụ và môi trường RStudio cho các nhà khoa học dữ liệu huấn luyện mô hình dự báo khiếu nại.")

    add_callout(
        doc,
        "Ý nghĩa kỹ thuật của Kiến trúc Lambda trong đề tài",
        "Kiến trúc Lambda giúp hệ thống giải quyết đồng thời hai bài toán tưởng chừng mâu thuẫn: xử lý hàng Terabyte dữ liệu lịch sử với độ chính xác tuyệt đối (thông qua Spark/MapReduce) nhưng vẫn duy trì khả năng cảnh báo sự cố viễn thông tức thời với độ trễ dưới 1 giây (thông qua HBase/Speed Layer)."
    )

print("Module 1 (Front, TOC, Intro, Chap 1) loaded successfully.")
